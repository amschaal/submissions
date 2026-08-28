"""
XLSX round-trip for the AgGrid table (list-of-objects) grid.

Two public entry points:

* ``build_table_template(schema, data)`` -> ``bytes`` of an ``.xlsx`` workbook
  populated with the current rows (if any) and with as much native Excel
  validation baked in as the format allows (dropdowns for ``enum``/boolean,
  decimal ranges for ``number``), plus column descriptions as cell notes.

* ``parse_table_xlsx(schema, file_obj)`` -> ``list[dict]`` of rows keyed by
  schema variable, ready to feed straight into
  :class:`dnaorder.validators.SamplesheetValidator`.

Excel's native validation is a UX aid only; the server-side validator remains
authoritative. Nested ``table``-type columns are intentionally skipped -- the
grid can't render them and they don't fit a flat sheet.

The template writes two header rows on the data sheet:

    row 1 (hidden) -- the machine variable id for each column
    row 2          -- the human title (frozen)
    row 3+         -- data

The hidden id row is what makes the round-trip robust: users can rename the
visible titles and reorder columns and we can still map each column back to the
right variable. :func:`parse_table_xlsx` falls back to matching visible titles
(or the raw variable id) when the hidden row has been removed.
"""
import csv
import io
import re

import xlsxwriter
from openpyxl import load_workbook

from dnaorder.spreadsheets import get_cols, get_data

DATA_SHEET = "Data"
INSTRUCTIONS_SHEET = "Instructions"
LISTS_SHEET = "_lists"

# Supported export/import formats. xlsx is the rich round-trip (validation,
# descriptions, dropdowns); csv/tsv are flat (variable-name header + data).
_DELIMITERS = {"csv": ",", "tsv": "\t"}
TABLE_FORMATS = {
    "xlsx": {
        "content_type":
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "extension": "xlsx",
    },
    "csv": {"content_type": "text/csv", "extension": "csv"},
    "tsv": {"content_type": "text/tab-separated-values", "extension": "tsv"},
}

ID_ROW = 0
TITLE_ROW = 1
DATA_START_ROW = 2

# Number of extra blank rows to pre-format so pasted data keeps the column
# validation/format. Kept modest; Excel applies column validation to the whole
# applied range regardless.
BLANK_ROWS = 500

# Values starting with any of these are treated as formulas by spreadsheet apps.
FORMULA_PREFIXES = ("=", "+", "-", "@")


def _norm(value):
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def _columns(schema):
    """Ordered list of ``(variable, property)`` for non-table columns."""
    return [(v, schema["properties"][v]) for v in get_cols(schema, table=False)]


def _title(variable, prop):
    return prop.get("title") or variable


def _is_required(schema, variable):
    return variable in schema.get("required", [])


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def build_table(schema, data=None, fmt="xlsx"):
    """Return the bytes of an exported table in ``fmt`` (xlsx / csv / tsv)."""
    if fmt == "xlsx":
        return build_table_template(schema, data)
    if fmt in _DELIMITERS:
        return build_table_delimited(schema, data, _DELIMITERS[fmt]).encode("utf-8")
    raise ValueError("Unsupported format: %s" % fmt)


def build_table_delimited(schema, data, delimiter):
    """Return CSV/TSV text: a variable-name header row followed by data rows.

    Flat by nature -- no descriptions, validation or examples (those are
    xlsx-only). Values are written verbatim (list values joined like the xlsx
    export); no spreadsheet formula-injection guard is applied, so the
    round-trip stays exact.
    """
    columns = _columns(schema)
    variables = [v for v, _ in columns]
    rows = get_data(
        {"order": variables, "properties": {v: p for v, p in columns}}, data
    )
    output = io.StringIO()
    writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
    writer.writerow(variables)
    if data:
        for row in rows:
            writer.writerow(["" if v is None else v for v in row])
    return output.getvalue()


def build_table_template(schema, data=None):
    """Return the bytes of an ``.xlsx`` template for ``schema``.

    ``data`` is an optional list of existing table rows (dicts keyed by
    variable) which are pre-populated into the sheet.
    """
    columns = _columns(schema)
    output = io.BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})

    title_fmt = workbook.add_format(
        {
            "bold": True,
            "bg_color": "#DDDDDD",
            "border": 1,
            "text_wrap": True,
            "valign": "top",
        }
    )
    required_title_fmt = workbook.add_format(
        {
            "bold": True,
            "bg_color": "#FBE4D5",
            "border": 1,
            "text_wrap": True,
            "valign": "top",
        }
    )
    # '@' forces Excel to treat the cell as text, preventing auto-coercion of
    # things like "MAR1" or "0012" into dates / numbers.
    text_fmt = workbook.add_format({"num_format": "@"})
    number_fmt = workbook.add_format({})

    sheet = workbook.add_worksheet(DATA_SHEET)
    lists = workbook.add_worksheet(LISTS_SHEET)
    lists.hide()

    sheet.set_row(ID_ROW, None, None, {"hidden": True})
    # Freeze below the title row so headers stay visible while scrolling.
    sheet.freeze_panes(DATA_START_ROW, 0)

    last_data_row = DATA_START_ROW + max(len(data or []), BLANK_ROWS) - 1
    list_col = 0

    for col, (variable, prop) in enumerate(columns):
        title = _title(variable, prop)
        required = _is_required(schema, variable)

        # Hidden machine-id row + visible title row.
        sheet.write_string(ID_ROW, col, variable)
        sheet.write_string(
            TITLE_ROW, col, title, required_title_fmt if required else title_fmt
        )

        description = prop.get("description")
        note = description or ""
        if required:
            note = ("Required. " + note).strip()
        if note:
            sheet.write_comment(
                TITLE_ROW, col, note, {"x_scale": 2, "y_scale": 2, "author": ""}
            )

        col_type = prop.get("type")
        col_fmt = number_fmt if col_type == "number" else text_fmt
        width = max(len(title) + 2, 14)
        sheet.set_column(col, col, width, col_fmt)

        _apply_validation(
            workbook, sheet, lists, col, prop, schema, variable, last_data_row,
            list_col=list_col,
        )
        if prop.get("enum") and not prop.get("multiple"):
            list_col += 1

    _write_data(sheet, columns, data, text_fmt, number_fmt)
    _write_instructions(workbook, schema, columns)

    workbook.close()
    output.seek(0)
    return output.getvalue()


def _apply_validation(
    workbook, sheet, lists, col, prop, schema, variable, last_row, list_col
):
    """Add native Excel data validation for a single column, when expressible."""
    first, last = DATA_START_ROW, last_row
    description = prop.get("description")
    required = _is_required(schema, variable)
    input_bits = []
    if required:
        input_bits.append("Required.")
    if description:
        input_bits.append(description)
    input_message = " ".join(input_bits)

    col_type = prop.get("type")
    enum = prop.get("enum")

    if enum and not prop.get("multiple"):
        # Write choices into the hidden list sheet and reference the range, so
        # we're not bound by the ~255 char inline-list limit.
        lc = list_col
        for row, choice in enumerate(enum):
            lists.write_string(row, lc, str(choice))
        col_letter = xlsxwriter.utility.xl_col_to_name(lc)
        source = "={sheet}!${c}$1:${c}${n}".format(
            sheet=LISTS_SHEET, c=col_letter, n=len(enum)
        )
        sheet.data_validation(
            first, col, last, col,
            _with_input({"validate": "list", "source": source}, input_message),
        )
    elif col_type == "boolean":
        sheet.data_validation(
            first, col, last, col,
            _with_input(
                {"validate": "list", "source": ["TRUE", "FALSE"]}, input_message
            ),
        )
    elif col_type == "number":
        rule = _number_rule(prop)
        sheet.data_validation(first, col, last, col, _with_input(rule, input_message))
    elif input_message:
        # No expressible constraint, but still surface the description/required
        # hint as an input message when the user selects the cell.
        sheet.data_validation(
            first, col, last, col,
            _with_input({"validate": "any"}, input_message),
        )


def _number_rule(prop):
    minimum = _to_number(prop.get("minimum"))
    maximum = _to_number(prop.get("maximum"))
    if minimum is not None and maximum is not None:
        return {"validate": "decimal", "criteria": "between",
                "minimum": minimum, "maximum": maximum}
    if minimum is not None:
        return {"validate": "decimal", "criteria": ">=", "value": minimum}
    if maximum is not None:
        return {"validate": "decimal", "criteria": "<=", "value": maximum}
    return {"validate": "decimal", "criteria": "between",
            "minimum": -1e300, "maximum": 1e300}


def _with_input(rule, message):
    if message:
        rule = dict(rule)
        rule["input_title"] = "Info"
        rule["input_message"] = message[:255]
        rule["show_input"] = True
    return rule


def _write_data(sheet, columns, data, text_fmt, number_fmt):
    if not data:
        return
    rows = get_data({"order": [v for v, _ in columns],
                     "properties": {v: p for v, p in columns}}, data)
    for r, row in enumerate(rows):
        for c, (variable, prop) in enumerate(columns):
            value = row[c] if c < len(row) else None
            _write_cell(sheet, DATA_START_ROW + r, c, value, prop, text_fmt, number_fmt)


def _write_cell(sheet, row, col, value, prop, text_fmt, number_fmt):
    if value is None or value == "":
        return
    if prop.get("type") == "number":
        number = _to_number(value)
        if number is not None:
            sheet.write_number(row, col, number, number_fmt)
            return
    # Always write text via write_string so a value like "=SUM(...)" is stored
    # as a literal string, never an executable formula (CSV/formula injection).
    sheet.write_string(row, col, _safe_text(value), text_fmt)


def _safe_text(value):
    text = value if isinstance(value, str) else str(value)
    return text


def _write_instructions(workbook, schema, columns):
    sheet = workbook.add_worksheet(INSTRUCTIONS_SHEET)
    header = workbook.add_format({"bold": True, "bg_color": "#DDDDDD", "border": 1})
    wrap = workbook.add_format({"text_wrap": True, "valign": "top"})
    sheet.set_column(0, 0, 24)
    sheet.set_column(1, 1, 12)
    sheet.set_column(2, 2, 10)
    sheet.set_column(3, 3, 60)
    sheet.write_row(0, 0, ["Column", "Type", "Required", "Description / rules"], header)
    for i, (variable, prop) in enumerate(columns):
        rules = _describe_rules(prop)
        desc = prop.get("description") or ""
        if rules:
            desc = (desc + "  " if desc else "") + rules
        sheet.write_row(
            i + 1, 0,
            [
                _title(variable, prop),
                prop.get("type", ""),
                "Yes" if _is_required(schema, variable) else "",
                desc,
            ],
            wrap,
        )


def _describe_rules(prop):
    bits = []
    if prop.get("unique"):
        bits.append("must be unique")
    if prop.get("enum") and not prop.get("multiple"):
        bits.append("one of: " + ", ".join(str(c) for c in prop["enum"]))
    if prop.get("enum") and prop.get("multiple"):
        bits.append("comma-separated, each one of: "
                    + ", ".join(str(c) for c in prop["enum"]))
    if prop.get("pattern"):
        bits.append("must match pattern %s" % prop["pattern"])
    mn, mx = _to_number(prop.get("minimum")), _to_number(prop.get("maximum"))
    if mn is not None and mx is not None:
        bits.append("between %s and %s" % (mn, mx))
    elif mn is not None:
        bits.append("minimum %s" % mn)
    elif mx is not None:
        bits.append("maximum %s" % mx)
    return "; ".join(bits)


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

def parse_table(schema, file_obj, fmt="xlsx"):
    """Parse an uploaded table file into a list of row dicts keyed by variable.

    ``fmt`` is one of ``xlsx`` / ``csv`` / ``tsv``.
    """
    if fmt == "xlsx":
        return parse_table_xlsx(schema, file_obj)
    if fmt in _DELIMITERS:
        return parse_table_delimited(schema, file_obj, _DELIMITERS[fmt])
    raise ValueError("Unsupported format: %s" % fmt)


def parse_table_xlsx(schema, file_obj, sheet_name=DATA_SHEET):
    """Parse an uploaded ``.xlsx`` into a list of row dicts keyed by variable.

    Column-to-variable mapping prefers the hidden machine-id row written by
    :func:`build_table_template`, and falls back to matching the visible title
    (or raw variable id) so hand-made or edited sheets still import.
    """
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames \
            else workbook.worksheets[0]
        rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_to_records(schema, rows)


def parse_table_delimited(schema, file_obj, delimiter):
    """Parse an uploaded CSV/TSV (variable-name header + data rows)."""
    text = file_obj.read()
    if isinstance(text, bytes):
        # utf-8-sig transparently strips a BOM that Excel may have written.
        text = text.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [tuple(r) for r in reader]
    return _rows_to_records(schema, rows)


def _rows_to_records(schema, rows):
    """Shared row-list -> record-list logic for every import format.

    Handles both the xlsx two-header-row layout (hidden id row + title row) and
    the single-header-row CSV/TSV layout: :func:`_find_header` locates the
    header, and a following row is only skipped when it, too, maps cleanly to
    the schema (i.e. it is the second header row, not data).
    """
    if not rows:
        return []

    header_index, col_map = _find_header(schema, rows)
    if col_map is None:
        raise ValueError(
            "Could not find a header row matching the schema. Expected columns "
            "such as: %s"
            % ", ".join(get_cols(schema, table=False)[:5])
        )

    data_start = header_index + 1
    if data_start < len(rows):
        follow = _match_row(schema, rows[data_start])
        if len(follow) >= len(col_map):
            data_start += 1

    records = []
    for row in rows[data_start:]:
        record = {}
        empty = True
        for idx, variable in col_map.items():
            value = row[idx] if idx < len(row) else None
            value = _coerce(value, schema["properties"][variable])
            if value not in (None, ""):
                empty = False
            record[variable] = value
        if not empty:
            records.append(record)
    return records


def _title_lookup(schema):
    """Map normalized title AND raw variable id -> variable."""
    lookup = {}
    for variable in get_cols(schema, table=False):
        prop = schema["properties"][variable]
        lookup[_norm(variable)] = variable
        lookup[_norm(_title(variable, prop))] = variable
    return lookup


def _match_row(schema, cells, lookup=None):
    lookup = lookup if lookup is not None else _title_lookup(schema)
    col_map = {}
    for idx, cell in enumerate(cells):
        if cell is None or str(cell).strip() == "":
            continue
        variable = lookup.get(_norm(cell))
        if variable is not None and idx not in col_map:
            col_map[idx] = variable
    return col_map


def _find_header(schema, rows, scan=5):
    lookup = _title_lookup(schema)
    best_index, best_map = None, {}
    for i, row in enumerate(rows[:scan]):
        col_map = _match_row(schema, row, lookup)
        if len(col_map) > len(best_map):
            best_index, best_map = i, col_map
    if not best_map:
        return None, None
    return best_index, best_map


def _coerce(value, prop):
    import datetime

    if value is None:
        return None
    col_type = prop.get("type")
    if isinstance(value, datetime.datetime):
        # Excel likely auto-converted a string into a date; best-effort recovery.
        value = value.date().isoformat() if value.time() == datetime.time(0, 0) \
            else value.isoformat()
    if col_type == "number":
        number = _to_number(value)
        return number if number is not None else str(value).strip()
    if col_type == "boolean":
        return _to_bool(value)
    if isinstance(value, float) and value.is_integer():
        # openpyxl returns floats for integer-looking cells; avoid "12.0".
        value = int(value)
    return str(value).strip()


def _to_number(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number.is_integer():
        return int(number)
    return number


def _to_bool(value):
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in ("true", "1", "yes", "y"):
        return True
    if text in ("false", "0", "no", "n", ""):
        return False
    return value
